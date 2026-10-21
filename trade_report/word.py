"""report.docx: the query written up as a short report with the charts and main tables."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import fmt
from .analysis import TOP_N, Analysis, growth, hhi_band
from .findings import (edition_summary, estimate_sentences, figures_note, findings,
                       gap_sentences)
from .model import period_label
from .summary import cell, kind_of

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def _shade(table_cell, colour: str):
    properties = table_cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), colour)
    properties.append(shading)


def _field(paragraph, code: str):
    """A Word field such as PAGE, which Word fills in when it shows the document."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"),
                       ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText" if text == code else "w:t")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    return run


class Report:
    def __init__(self, a: Analysis, charts: list[Path], files: dict[str, str]):
        self.a, self.q, self.charts, self.files = a, a.q, {c.stem: c for c in charts}, files
        self.doc = Document()
        self.number = 0
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2))
        styles = self.doc.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(10.5)
        for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
            styles[name].font.name = "Calibri"
            styles[name].font.size = Pt(size)
            styles[name].font.color.rgb = NAVY
            styles[name].font.bold = True
        footer = section.footer.paragraphs[0]
        footer.text = self.q.source_note
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = GREY
        # Page numbers: 'Page 3 of 9'
        numbers = section.footer.add_paragraph()
        numbers.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        runs = [numbers.add_run("Page "), _field(numbers, "PAGE"), numbers.add_run(" of "),
                _field(numbers, "NUMPAGES")]
        for run in runs:
            run.font.size = Pt(8)
            run.font.color.rgb = GREY

    # -- building blocks ---------------------------------------------------------------
    def heading(self, text):
        """A numbered section; every section after the first starts on a new page."""
        self.number += 1
        heading = self.doc.add_heading(f"{self.number}. {text}", level=1)
        heading.paragraph_format.page_break_before = self.number > 1

    def paragraph(self, text, italic=False, size=None, colour=None):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if colour:
            run.font.color.rgb = colour
        return para

    def bullets(self, items):
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def table(self, headers, rows, numeric):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, header in enumerate(headers):
            cell_ = table.rows[0].cells[j]
            cell_.text = ""
            run = cell_.paragraphs[0].add_run(header)
            run.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell_, "1F4E79")
        for row in rows:
            cells = table.add_row().cells
            for j, value in enumerate(row):
                cells[j].text = ""
                para = cells[j].paragraphs[0]
                run = para.add_run(str(value))
                run.font.size = Pt(8.5)
                if value in ("Missing", "Mixed units"):
                    run.italic = True
                    run.font.color.rgb = RGBColor(0xC5, 0x5A, 0x11)
                if numeric[j]:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.doc.add_paragraph()

    def chart(self, stem, caption):
        path = self.charts.get(stem)
        if not path:
            return False
        self.doc.add_picture(str(path), width=Cm(17))
        self.paragraph(caption, italic=True, size=9, colour=GREY)
        return True

    # -- the report --------------------------------------------------------------------
    def build(self, path: Path) -> Path:
        a, q, s, m = self.a, self.q, self.a.summary, self.a.metric
        product = (fmt.lower_first(q.product_desc)
                   if q.product_code not in ("TOTAL", "AG2", "AG4", "AG6") else "all products")
        self.doc.add_paragraph(f"{q.reporter_name}'s {q.flow.lower()} of {product}, {q.span}",
                               style="Title")
        self.paragraph(f"{q.product_label} · {q.source} · generated {dt.date.today():%d %B %Y} "
                       "by the International Trade Analyzer", italic=True, colour=GREY)

        self.heading("Summary")
        self.bullets(findings(a))
        figures = [(f"Total {m.label.lower()}", fmt.amount(s["total"], m))]
        if len(a.years) >= 2:
            figures += [(f"CAGR {a.years[0]}–{a.years[-1]}", fmt.pct(s["cagr"], signed=True)),
                        ("Trend growth", fmt.pct(s["trend"], signed=True))]
            if "real_cagr" in s:
                figures.append((f"CAGR in constant {a.real_base} USD",
                                fmt.pct(s["real_cagr"], signed=True)))
        for metric in a.metrics[1:]:
            figures.append((f"Total {metric.label.lower()}",
                            fmt.amount(s["totals"][metric.key], metric)))
        figures += [(q.roles.capitalize(), str(s["partners"])),
                    ("Concentration (HHI)", f"{fmt.count(s['hhi'])} {hhi_band(s['hhi'])}".strip())]
        self.table(["Measure", "Figure"], [list(f) for f in figures], [False, True])

        self.heading("Scope, data and method")
        self.bullets([
            f"Source: {q.source_description or q.source}.",
            f"Reporter: {q.reporter_name}; flow: {q.flow.lower()}; partners are {q.roles}.",
            f"Product: {q.product_label}; one product row is one {q.product_level}.",
            f"Period: {q.span} ({'monthly' if q.monthly else 'annual'}).",
            f"HS edition: {edition_summary(a)}.",
            figures_note(a),
            f"Measured by {m.heading}" + (f"; also reported: "
                                          f"{', '.join(x.heading for x in a.metrics[1:])}."
                                          if len(a.metrics) > 1 else "."),
        ])
        self.paragraph("Missing data. Figures are shown as the source reports them. Where the "
                       "source has no value the report says Missing, and so does every total, "
                       "share, unit value and growth rate that depends on it. Nothing is "
                       "estimated, interpolated or filled in.")
        if q.missing_rules:
            self.bullets(q.missing_rules)
        if q.estimate_rules:
            self.paragraph(f"Estimates in the source. A figure {q.source} itself marks as an "
                           f"estimate is part of the database: it is used as {q.source} gives it "
                           "and stated as an estimate, in this report and in the workbook (in "
                           "blue, and in the Source estimates sheet). The program itself never "
                           "estimates anything.")
            self.bullets(q.estimate_rules)
        self.paragraph("CAGR is (last year / first year)^(1 / years between) − 1; trend growth is "
                       "the annual growth of a log-linear trend fitted to every year. HHI is the "
                       "sum of squared shares x 10,000 (above 2,500: highly concentrated)."
                       + (f" Constant dollars use the US consumer price index ({a.cpi_note})."
                          if a.real_totals is not None else ""), size=9.5)

        if self.charts.get("01_overview"):
            self.heading("The trend")
            self.chart("01_overview", f"{m.label} by {a.period_word}, with the headline figures.")
            self.period_table()
        if self.charts.get("02_partners") or len(a.partners):
            self.heading(q.roles.capitalize())
            self.chart("02_partners", f"{m.label} by {q.role} over {q.span}.")
            self.top_table(a.partners, q.role.capitalize())
        self.chart_section("03_partners_by_period",
                           f"{a.period_word.capitalize()} by {a.period_word}",
                           f"{m.label} by {q.role} and {a.period_word}.")
        self.chart_section("04_shares", "Shares", f"Share of the {m.label.lower()} by {q.role}.")
        self.chart_section("05_growth", "Growth", "Compound annual growth and change in share.")
        self.chart_section("06_concentration", "Concentration",
                           f"HHI and the share of the largest {q.roles}.")
        if self.chart_section("07_products", "Products", f"{m.label} by product."):
            self.top_table(a.products, "Product")
        if a.focus is not None and self.chart_section(
                "08_partner_focus", a.partner_names[q.partner_code],
                f"{a.partner_names[q.partner_code]} against all other {q.roles}."):
            self.focus_table()
        self.chart_section("09_seasonality", "Seasonality", f"{m.label} by month and year.")

        estimates = estimate_sentences(a)
        self.heading("Data gaps, source estimates and checks" if estimates else
                     "Data gaps and checks")
        self.bullets(gap_sentences(a) or ["No figure used in this report is missing from the "
                                          "source."])
        self.bullets(estimates)
        self.bullets([f"{check}: {result}" for check, result in q.checks])
        self.paragraph("Files: " + "; ".join(f"{k}: {v}" for k, v in self.files.items()),
                       size=9, colour=GREY)
        self.doc.save(str(path))
        return path

    def chart_section(self, stem, title, caption) -> bool:
        if not self.charts.get(stem):
            return False
        self.heading(title)
        return self.chart(stem, caption)

    def period_table(self):
        a, q, m = self.a, self.q, self.a.metric
        if len(a.periods) > 24:
            self.paragraph(f"The {a.period_word}-by-{a.period_word} table is in the workbook.",
                           italic=True, size=9, colour=GREY)
            return
        totals = a.period_totals[m.key]
        estimated = [x for x in a.metrics if x.key in a.estimated]
        headers = ["Month" if q.monthly else "Year", m.heading, "Growth"]
        headers += [x.heading for x in a.metrics[1:]]
        headers += [f"{x.label}: estimated by {q.estimated_by}" for x in estimated]
        headers += ["Largest share", "HHI"]
        rows = []
        for i, period in enumerate(a.periods):
            before = a.periods[i - 1] if i else None
            row = [period_label(period, q.monthly, short=True),
                   cell(totals[period], kind_of(m), m),
                   fmt.pct(growth(totals[period], totals[before]), 1, signed=True)
                   if before else "n/a"]
            row += [cell(a.period_totals[x.key][period], kind_of(x), x) for x in a.metrics[1:]]
            row += [cell(a.estimated[x.key].by_period[period], "share") for x in estimated]
            stats = a.period_stats.loc[period]
            row += [cell(stats["Largest share"], "share"), cell(stats["HHI"], "count")]
            rows.append(row)
        self.table(headers, rows, [False] + [True] * (len(headers) - 1))

    def top_table(self, frame, label):
        m = self.a.metric
        top = frame.head(TOP_N)
        rows = [[str(r["Rank"]), str(r["Name"])[:60], cell(r[m.heading], kind_of(m), m),
                 cell(r["Share"], "share"), cell(r.get("CAGR", float("nan")), "growth")]
                for r in top.to_dict("records")]
        self.table(["Rank", label, m.heading, "Share", "CAGR"], rows,
                   [True, False, True, True, True])

    def focus_table(self):
        a, q = self.a, self.q
        f = a.focus
        if len(f) > 24:
            return
        name = a.partner_names[q.partner_code]
        mine = f"{name}: {a.metric.heading}"
        rows = [[period_label(r["Period"], q.monthly, short=True),
                 cell(r[mine], kind_of(a.metric), a.metric), cell(r["Share"], "share"),
                 str(r["Rank"]), fmt.pct(r[f"{name}: growth"], 1, signed=True)]
                for r in f.to_dict("records")]
        self.table(["Month" if q.monthly else "Year", a.metric.heading, "Share", "Rank",
                    "Growth"], rows, [False, True, True, True, True])


def write_report(a: Analysis, charts: list[Path], path: Path, files: dict[str, str]) -> Path:
    return Report(a, charts, files).build(path)
